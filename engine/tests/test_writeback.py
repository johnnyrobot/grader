"""Annotated write-back: the copy gets the comments, the original never changes.

"NEVER mutate the original" is the first line of writeback.py and was asserted
nowhere. It is the one property a teacher cannot check by eye — they open the
annotated copy, not the file the student submitted — so it needs a test per
format rather than a comment.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parents[1]))

from grader.writeback import (  # noqa: E402  (needs the path insert above)
    Comment,
    annotate,
    annotate_docx,
    annotate_pdf,
    annotate_xlsx,
    feedback_report,
)


# ---------------------------------------------------------------------------
# Fixtures: real files, because the contract is about bytes on disk
# ---------------------------------------------------------------------------

@pytest.fixture
def docx_file(tmp_path):
    import docx

    path = tmp_path / "submission.docx"
    doc = docx.Document()
    doc.add_paragraph("The thesis is stated in the opening paragraph.")
    doc.add_paragraph("Evidence follows in the second paragraph.")
    doc.save(str(path))
    return path


@pytest.fixture
def pdf_file(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "submission.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    with open(path, "wb") as f:
        writer.write(f)
    return path


@pytest.fixture
def xlsx_file(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "submission.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Question 1"
    wb.active["B3"] = "42"
    wb.save(str(path))
    return path


def _unchanged(path):
    """A closure over the file's bytes, asked again after annotation."""
    before = path.read_bytes()
    return lambda: path.read_bytes() == before


# ---------------------------------------------------------------------------
# The contract, per format
# ---------------------------------------------------------------------------

def test_docx_annotation_leaves_the_original_untouched(docx_file, tmp_path):
    still_the_same = _unchanged(docx_file)

    result = annotate_docx(docx_file, tmp_path / "annotated.docx",
                           [Comment(text="Sharpen this", anchor_text="thesis")])

    assert still_the_same(), "the student's own file was modified"
    assert Path(result.output_path).exists()
    assert result.comments_added == 1


def test_pdf_annotation_leaves_the_original_untouched(pdf_file, tmp_path):
    still_the_same = _unchanged(pdf_file)

    result = annotate_pdf(pdf_file, tmp_path / "annotated.pdf",
                          [Comment(text="Check this figure", page=1)])

    assert still_the_same(), "the student's own file was modified"
    assert result.comments_added == 1


def test_xlsx_annotation_leaves_the_original_untouched(xlsx_file, tmp_path):
    still_the_same = _unchanged(xlsx_file)

    result = annotate_xlsx(xlsx_file, tmp_path / "annotated.xlsx",
                           [Comment(text="Wrong formula", cell_ref="B3")])

    assert still_the_same(), "the student's own file was modified"
    assert result.comments_added == 1


def test_the_dispatcher_never_writes_over_the_source(docx_file, tmp_path):
    still_the_same = _unchanged(docx_file)

    annotate(docx_file, tmp_path / "out.docx", [Comment(text="Good")])

    assert still_the_same()


# ---------------------------------------------------------------------------
# The comments actually land
# ---------------------------------------------------------------------------

def test_a_pdf_comment_is_written_onto_the_page_it_names(pdf_file, tmp_path):
    """Annotating page 2's problem onto page 1 sends the student to the wrong place."""
    from pypdf import PdfReader

    dest = tmp_path / "annotated.pdf"
    annotate_pdf(pdf_file, dest, [Comment(text="Second page issue", page=2)])

    pages = PdfReader(str(dest)).pages
    assert "/Annots" not in pages[0] or not pages[0]["/Annots"]
    assert pages[1]["/Annots"]


def test_a_pdf_comment_past_the_last_page_lands_on_the_last_page(pdf_file, tmp_path):
    """Better the last page than an IndexError halfway through a batch."""
    from pypdf import PdfReader

    dest = tmp_path / "annotated.pdf"
    result = annotate_pdf(pdf_file, dest, [Comment(text="Overflow", page=99)])

    assert result.comments_added == 1
    assert PdfReader(str(dest)).pages[-1]["/Annots"]


def test_an_xlsx_comment_lands_on_the_named_cell(xlsx_file, tmp_path):
    from openpyxl import load_workbook

    dest = tmp_path / "annotated.xlsx"
    annotate_xlsx(xlsx_file, dest, [Comment(text="Wrong formula", cell_ref="B3")])

    sheet = load_workbook(str(dest)).active
    assert "Wrong formula" in sheet["B3"].comment.text
    assert sheet["A1"].comment is None


def test_an_xlsx_comment_without_a_cell_falls_back_to_a1(xlsx_file, tmp_path):
    from openpyxl import load_workbook

    dest = tmp_path / "annotated.xlsx"
    annotate_xlsx(xlsx_file, dest, [Comment(text="General note")])

    assert "General note" in load_workbook(str(dest)).active["A1"].comment.text


def test_a_docx_comment_with_no_anchor_still_gets_recorded(docx_file, tmp_path):
    result = annotate_docx(docx_file, tmp_path / "out.docx",
                           [Comment(text="Overall: solid work")])

    assert result.comments_added == 1


def test_a_docx_anchor_that_is_not_in_the_document_is_not_silently_dropped(
    docx_file, tmp_path
):
    """A comment the teacher wrote must reach the student somewhere."""
    result = annotate_docx(docx_file, tmp_path / "out.docx",
                           [Comment(text="On the conclusion", anchor_text="nowhere")])

    assert result.comments_added == 1


def test_every_comment_is_accounted_for(docx_file, tmp_path):
    comments = [
        Comment(text="One", anchor_text="thesis"),
        Comment(text="Two", anchor_text="Evidence"),
        Comment(text="Three"),
    ]

    result = annotate_docx(docx_file, tmp_path / "out.docx", comments)

    assert result.comments_added == len(comments)


# ---------------------------------------------------------------------------
# Formats that cannot carry comments
# ---------------------------------------------------------------------------

def test_an_unsupported_format_gets_a_feedback_report_instead(tmp_path):
    """A .txt submission still has to produce something the teacher can hand back."""
    src = tmp_path / "submission.txt"
    src.write_text("plain text submission")
    still_the_same = _unchanged(src)

    result = annotate(src, tmp_path / "feedback.docx",
                      [Comment(text="Expand the argument")])

    assert still_the_same()
    assert Path(result.output_path).exists()


def test_the_feedback_report_carries_the_scores_and_the_words(tmp_path):
    import docx

    dest = tmp_path / "nested" / "feedback.docx"
    feedback_report(
        dest,
        draft_text="",
        rubric_title="Essay Rubric",
        scores={"thesis": 8.0, "evidence": 4.5},
        feedback="Strong thesis, thin evidence.",
        total=12.5,
        max_total=15.0,
    )

    text = "\n".join(p.text for p in docx.Document(str(dest)).paragraphs)
    cells = [c.text for t in docx.Document(str(dest)).tables for r in t.rows for c in r.cells]

    assert "Essay Rubric" in text
    assert "Strong thesis, thin evidence." in text
    assert "thesis" in cells and "8" in cells
    assert "12.5" in cells and "15" in cells


def test_the_feedback_report_creates_its_parent_directory(tmp_path):
    dest = tmp_path / "does" / "not" / "exist" / "feedback.docx"

    feedback_report(dest, draft_text="", rubric_title="R", scores={},
                    feedback="", total=0, max_total=0)

    assert dest.exists()
